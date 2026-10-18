"""Economic units, causal intervals and report selection for the bridge."""
import importlib.util
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd

MODULE = Path(__file__).resolve().parents[1]/"scripts/build_report.py"
spec = importlib.util.spec_from_file_location("build_report", MODULE)
report = importlib.util.module_from_spec(spec)
spec.loader.exec_module(report)


class ReportTests(unittest.TestCase):
    def test_exact_growth_and_level(self):
        levels = pd.Series(100., index=pd.period_range("2023Q1", "2024Q1", freq="Q"))
        result = report.measures(.01, pd.Period("2024Q2"), levels)
        self.assertAlmostEqual(result["qoq_percent"], 100*np.expm1(.01))
        self.assertAlmostEqual(result["level_billions"], .1*np.exp(.01))
        self.assertAlmostEqual(result["yoy_percent"], 100*np.expm1(.01))
        self.assertNotEqual(result["qoq_percent"], 1.)

    def test_carry_over_for_each_quarter(self):
        levels = pd.Series([100.]*4+[101., 102., 103.],
                           index=pd.period_range("2023Q1", periods=7, freq="Q"))
        for quarter, expected in [(1, 2.), (2, 1.75), (4, 2.5)]:
            q = pd.Period(f"2024Q{quarter}")
            target = 102. if quarter != 4 else 104.
            value = report.measures(np.log(target/levels.loc[q-1]), q, levels)
            self.assertAlmostEqual(value["carry_over_percent"], expected)

    def test_missing_carry_input_is_not_zero(self):
        levels = pd.Series([100.]*5, index=pd.period_range("2023Q1", periods=5, freq="Q"))
        levels.loc["2023Q2"] = np.nan
        self.assertTrue(np.isnan(report.measures(.01, pd.Period("2024Q2"), levels)["carry_over_percent"]))

    def test_bands_use_same_stage_and_only_released_gdp(self):
        dates = pd.date_range("2000-01-01", periods=15, freq="MS")
        forecasts = pd.Series(np.arange(15)/100, index=dates)
        ends = dates[dates.month % 3 == 0]
        growth = pd.Series([1., 2., 3., 4., 5.], index=ends)
        bands = report.error_bands(forecasts, growth, minimum=4)
        self.assertTrue(bands.lower.loc["2000-12-01"] != bands.lower.loc["2000-12-01"])
        self.assertEqual(bands.n_errors.loc["2001-01-01"], 4)
        past = growth.iloc[:4].to_numpy()-forecasts.iloc[[0, 3, 6, 9]].to_numpy()
        self.assertAlmostEqual(bands.lower.loc["2001-01-01"], forecasts.iloc[12]+np.quantile(past, .05))
        changed = growth.copy()
        changed.iloc[-1] = -900
        pd.testing.assert_frame_equal(report.error_bands(forecasts, changed, minimum=4), bands)

    def test_level_anchor_matches_each_months_target_quarter(self):
        dates = pd.date_range("2024-01-01", periods=6, freq="MS")
        levels = pd.Series([100., 102.], index=pd.period_range("2023Q4", periods=2, freq="Q"))
        anchors = levels.reindex(dates.to_period("Q")-1).to_numpy()
        np.testing.assert_array_equal(anchors, [100., 100., 100., 102., 102., 102.])

    def test_score_excludes_unreleased_quarter(self):
        dates = pd.date_range("2007-04-01", periods=9, freq="MS")
        history = pd.DataFrame({"nowcast": .01, "actual_log": 0., "naive_log": .02,
                                "lower_log": -.01, "upper_log": .03}, index=dates)
        scores = report.score_history(history, pd.Timestamp("2007-12-01"))
        row = scores[(scores.period == "All") & (scores.stage == 0)].iloc[0]
        self.assertEqual(row.n, 6)
        self.assertAlmostEqual(row.rmse, 1.)
        self.assertAlmostEqual(row.relative_naive_rmse, .5)

    def test_live_boundary_uses_reference_month_and_labels_release_month(self):
        dates = pd.date_range("2026-08-01", periods=2, freq="MS")
        actual = pd.Series([.2], index=dates[-1:])
        history = pd.DataFrame({"growth": [.1, .3], "growth_lower": [-.1, .1],
                                "growth_upper": [.3, .5], "level": [100., 101.],
                                "level_lower": [99., 100.], "level_upper": [101., 102.]}, index=dates)
        for is_level in [False, True]:
            chart = report.figure(history, actual, is_level)
            boundary = next(shape for shape in chart.layout.shapes if shape.xref == "x")
            self.assertEqual(pd.Timestamp(boundary.x0), pd.Timestamp("2026-10-01"))
            self.assertIn("Nov 2026", chart.layout.annotations[0].text)
            column = "level" if is_level else "growth"
            np.testing.assert_array_equal(chart.data[2].y, history[column])
            self.assertEqual(chart.data[2].name, "Pseudo-real-time nowcast")

    def test_report_runs_without_local_notes(self):
        with TemporaryDirectory() as folder:
            base = Path(folder)
            for subfolder in ["HISTORY/PREDICTION", "DATA/GDP_History", "DATA/VINTAGE_METADATA"]:
                (base/subfolder).mkdir(parents=True)
            dates = pd.date_range("2005-01-01", "2021-09-01", freq="MS")
            q = dates.month % 3 == 0
            levels = np.full(len(dates), np.nan)
            levels[q] = 100*np.exp(.002*np.arange(q.sum()))
            growth = np.where(q, .002, np.nan)
            pd.DataFrame({"Level": levels, "Growth": growth}, index=dates).to_csv(
                base/"DATA/GDP_History/GDP20211009.csv", index_label="Time")
            origins = pd.date_range("2006-04-01", "2021-09-01", freq="MS")
            pd.DataFrame({"nowcast": .002}, index=origins).to_csv(
                base/"HISTORY/PREDICTION/20211009_bridge.csv", index_label="Time")
            metadata = {"vintage": "20211009", "eval_end": "2021-09-01",
                        "produced_at": "2021-10-09T12:00:00+02:00",
                        "model": "ip_eurocoin_pca_2", "model_version": "test"}
            (base/"HISTORY/PREDICTION/20211009_bridge.json").write_text(json.dumps(metadata))
            (base/"DATA/VINTAGE_METADATA/vintage20211009.json").write_text(json.dumps(
                {"gdp": {"url": "https://example.test/gdp"}}))
            assets = report.build_report(base, "20211009")
            latest = json.loads((assets/"latest_itnow.json").read_text())
            self.assertEqual(latest["last_actual_quarter"], "2021Q2")
            self.assertAlmostEqual(latest["current"]["qoq_percent"], 100*np.expm1(.002))
            self.assertTrue((assets/"GDP_Level_Ncast.json").is_file())
            self.assertFalse((assets/"MODEL.pdf").exists())
            self.assertFalse((base/"HISTORY/RELEASES").exists())

    def test_staging_preserves_release_workbooks(self):
        with TemporaryDirectory() as folder:
            root = Path(folder)
            website, assets, project = root/"site", root/"assets", root/"project"
            destination = website/"static/ITNOW"
            destination.mkdir(parents=True)
            (website/"content").mkdir()
            assets.mkdir()
            (project/"website").mkdir(parents=True)
            for name in ["MODEL.pdf", "GDP_MEASURES.pdf"]:
                (destination/name).write_bytes(b"method note")
            names = ["ITNOW_History.xlsx", "ITNOW_History_pre_bridge.xlsx"]
            for name in names:
                (destination/name).write_bytes(b"published records")
                (assets/name).write_bytes(b"retrospective report")
            (assets/"release_record.csv").write_text("status\nunpublished\n")
            (assets/"latest_itnow.json").write_text("{}")
            report.stage_website(website, assets, project)
            for name in names:
                self.assertEqual((destination/name).read_bytes(), b"published records")
            self.assertFalse((destination/"release_record.csv").exists())
            self.assertTrue((destination/"latest_itnow.json").is_file())


class LiveHistoryTests(unittest.TestCase):
    def setUp(self):
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill

        self.temporary = TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        root = Path(self.temporary.name)
        self.base, self.website = root/"base", root/"site"
        self.assets, self.project = root/"assets", root/"project"
        self.local = self.base/"HISTORY/RELEASES/ITNOW_History.xlsx"
        self.public = self.website/"static/ITNOW/ITNOW_History.xlsx"
        self.local.parent.mkdir(parents=True)
        self.public.parent.mkdir(parents=True)
        (self.website/"content").mkdir()
        (self.project/"website").mkdir(parents=True)
        self.assets.mkdir()
        for name in ["MODEL.pdf", "GDP_MEASURES.pdf"]:
            (self.assets/name).write_bytes(b"method note")
        workbook = Workbook()
        workbook.active.title = "INFO"
        workbook["INFO"]["A1"] = "Release definitions"
        sheet = workbook.create_sheet("NOWCAST")
        sheet.append(["Release date", "Reference month", "Target quarter", "QoQ", "Lower 5%",
                      "Upper 95%", "YoY", "Carry-over", "GDP level (EUR bn)", "Data vintage",
                      "Produced at", "Code version"])
        for column, width in [("A", 20), ("B", 22), ("K", 32)]:
            sheet.column_dimensions[column].width = width
        for column in range(1, 13):
            sheet.cell(1, column).fill = PatternFill("solid", fgColor="155E63")
            sheet.cell(2, column).number_format = (
                "0.00%" if 4 <= column <= 8 else "dd mmm yyyy" if column in [1, 2, 10] else "General")
        sheet.freeze_panes = "C2"
        sheet.auto_filter.ref = "A1:L101"
        sheet.cell(101, 1).number_format = "dd mmm yyyy"
        workbook.create_sheet("DATA")["A1"] = "Indicator definitions"
        workbook.save(self.local)
        self.original = self.local.read_bytes()
        self.public.write_bytes(self.original)
        self.latest = {"reference_date": "2026-10-01", "produced_at": "2026-11-09T12:00:00+01:00",
                       "data_vintage": "20261109", "quarter": "Q4 2026", "code_version": "test-1",
                       "current": {"qoq_percent": .51, "lower_percent": -.21, "upper_percent": 1.42,
                                   "yoy_percent": 1.31, "carry_over_percent": 1.04,
                                   "level_billions": 496.60}}

    def prepare(self, date="2026-11-09"):
        return report.live_history(self.base, self.website, self.latest, pd.Timestamp(date))

    def stage(self, date="2026-11-09"):
        report.stage_website(self.website, self.assets, self.project, self.prepare(date))

    def test_live_release_is_copied_with_correct_units_and_sheets(self):
        from openpyxl import load_workbook

        self.stage()
        self.assertEqual(self.local.read_bytes(), self.public.read_bytes())
        workbook = load_workbook(self.local)
        self.assertEqual(workbook.sheetnames, ["INFO", "NOWCAST", "DATA"])
        self.assertEqual(workbook["INFO"]["A1"].value, "Release definitions")
        self.assertEqual(workbook["DATA"]["A1"].value, "Indicator definitions")
        sheet = workbook["NOWCAST"]
        self.assertEqual(sheet["A2"].value, pd.Timestamp("2026-11-09").to_pydatetime())
        self.assertEqual(sheet["B2"].value, pd.Timestamp("2026-10-01").to_pydatetime())
        np.testing.assert_allclose([sheet.cell(2, c).value for c in range(4, 9)],
                                   [.0051, -.0021, .0142, .0131, .0104])
        self.assertEqual(sheet["I2"].value, 496.60)
        self.assertEqual(sheet["J2"].value, pd.Timestamp("2026-11-09").to_pydatetime())
        self.assertEqual(sheet["K2"].value, self.latest["produced_at"])
        self.assertEqual(sheet["L2"].value, "test-1")
        self.assertEqual(sheet["D2"].number_format, "0.00%")
        self.assertEqual(sheet.freeze_panes, "C2")
        self.assertEqual(sheet.auto_filter.ref, "A1:L101")
        self.assertEqual(sheet.column_dimensions["K"].width, 32)
        self.assertEqual(sheet["A1"].fill.fgColor.rgb, "00155E63")

    def test_prelaunch_historical_and_future_runs_do_not_enter_history(self):
        cases = [
            ("2026-10-18", "2026-09-01", "2026-10-17T12:00:00+02:00", "20261009"),
            ("2026-11-09", "2026-09-01", "2026-11-09T12:00:00+01:00", "20261109"),
            ("2026-12-09", "2026-10-01", "2026-11-09T12:00:00+01:00", "20261109"),
            ("2026-11-09", "2026-10-01", "2026-10-31T12:00:00+01:00", "20261031"),
            ("2026-11-09", "2026-10-01", "2026-11-10T12:00:00+01:00", "20261109"),
        ]
        for date, reference, produced, vintage in cases:
            with self.subTest(date=date, reference=reference, produced=produced):
                self.latest.update(reference_date=reference, produced_at=produced, data_vintage=vintage)
                self.assertIsNone(self.prepare(date))
        self.assertEqual(self.local.read_bytes(), self.original)
        self.assertEqual(self.public.read_bytes(), self.original)

    def test_repeating_a_report_does_not_duplicate_or_redate_release(self):
        self.stage()
        recorded = self.local.read_bytes()
        self.stage("2026-11-10")
        self.assertEqual(self.local.read_bytes(), recorded)
        self.assertEqual(self.public.read_bytes(), recorded)
        self.latest["current"]["qoq_percent"] = 99
        with self.assertRaisesRegex(ValueError, "will not be overwritten"):
            self.prepare("2026-11-10")
        self.assertEqual(self.local.read_bytes(), recorded)

    def test_missing_bounds_remain_blank_and_zero_growth_is_preserved(self):
        from openpyxl import load_workbook

        self.latest["current"].update(lower_percent=None, upper_percent=np.nan, qoq_percent=0)
        self.stage()
        sheet = load_workbook(self.local)["NOWCAST"]
        self.assertEqual(sheet["D2"].value, 0)
        self.assertIsNone(sheet["E2"].value)
        self.assertIsNone(sheet["F2"].value)

    def test_a_new_release_preserves_earlier_rows(self):
        from openpyxl import load_workbook

        self.stage()
        first = list(load_workbook(self.local)["NOWCAST"].values)[1]
        self.latest.update(produced_at="2026-12-09T12:00:00+01:00", data_vintage="20261209",
                           reference_date="2026-11-01")
        self.latest["current"]["qoq_percent"] = .7
        self.stage("2026-12-09")
        rows = list(load_workbook(self.local)["NOWCAST"].values)
        self.assertEqual(rows[1], first)
        self.assertEqual(rows[2][1], pd.Timestamp("2026-11-01").to_pydatetime())
        self.assertAlmostEqual(rows[2][3], .007)
        self.assertEqual(rows[2][0], pd.Timestamp("2026-12-09").to_pydatetime())

    def test_website_history_can_initialize_a_missing_local_copy(self):
        self.local.unlink()
        self.stage()
        self.assertEqual(self.local.read_bytes(), self.public.read_bytes())

    def test_full_report_records_a_live_release_automatically(self):
        from openpyxl import load_workbook

        for folder in ["HISTORY/PREDICTION", "DATA/GDP_History", "DATA/VINTAGE_METADATA"]:
            (self.base/folder).mkdir(parents=True, exist_ok=True)
        dates = pd.date_range("2005-01-01", "2026-10-01", freq="MS")
        quarter_end = dates.month % 3 == 0
        levels = np.full(len(dates), np.nan)
        levels[quarter_end] = 100*np.exp(.002*np.arange(quarter_end.sum()))
        pd.DataFrame({"Level": levels, "Growth": np.where(quarter_end, .002, np.nan)},
                     index=dates).to_csv(self.base/"DATA/GDP_History/GDP20261109.csv", index_label="Time")
        origins = dates[dates >= "2006-04-01"]
        pd.DataFrame({"nowcast": .003}, index=origins).to_csv(
            self.base/"HISTORY/PREDICTION/20261109_bridge.csv", index_label="Time")
        metadata = {"vintage": "20261109", "eval_end": "2026-10-01",
                    "produced_at": self.latest["produced_at"], "model": "ip_eurocoin_pca_2",
                    "model_version": "test", "code_version": "test-1"}
        (self.base/"HISTORY/PREDICTION/20261109_bridge.json").write_text(json.dumps(metadata))
        (self.base/"DATA/VINTAGE_METADATA/vintage20261109.json").write_text(
            json.dumps({"gdp": {"url": "https://example.test/gdp"}}))
        for name in ["MODEL.pdf", "GDP_MEASURES.pdf"]:
            (self.public.parent/name).write_bytes(b"method note")
        prepare = report.live_history
        with patch.object(report, "live_history", side_effect=lambda base, site, latest, date:
                          prepare(base, site, latest, pd.Timestamp("2026-11-09"))):
            report.build_report(self.base, "20261109", self.website)
        row = list(load_workbook(self.local)["NOWCAST"].values)[1]
        self.assertAlmostEqual(row[3], np.expm1(.003))
        self.assertEqual(row[1], pd.Timestamp("2026-10-01").to_pydatetime())
        self.assertEqual(self.local.read_bytes(), self.public.read_bytes())

    def test_published_rows_cannot_be_lost_to_a_stale_local_workbook(self):
        self.stage()
        published = self.public.read_bytes()
        self.local.write_bytes(self.original)
        with self.assertRaisesRegex(ValueError, "histories differ"):
            self.prepare()
        self.assertEqual(self.public.read_bytes(), published)

    def test_failed_workbook_copy_restores_both_copies(self):
        original_write = Path.write_bytes
        def fail_once(path, content):
            if path == self.public and content != self.original:
                raise OSError("Synthetic website write failure")
            return original_write(path, content)
        release = self.prepare()
        with patch.object(Path, "write_bytes", fail_once):
            with self.assertRaisesRegex(OSError, "Synthetic website"):
                report.stage_website(self.website, self.assets, self.project, release)
        self.assertEqual(self.local.read_bytes(), self.original)
        self.assertEqual(self.public.read_bytes(), self.original)


if __name__ == "__main__":
    unittest.main()

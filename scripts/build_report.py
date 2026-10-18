#!/usr/bin/env python3
"""Build bridge figures, GDP measures and diagnostics from decimal log growth."""

from __future__ import annotations

import argparse
from copy import copy
import hashlib
from io import BytesIO
import json
import shutil
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.graph_objects as go
from plotly.offline import get_plotlyjs

FIRST_LIVE_RELEASE = pd.Timestamp("2026-11-01")
ONLINE_START = FIRST_LIVE_RELEASE - pd.DateOffset(months=1)
SCORE_START = pd.Period("2007Q2")
CRISES = [("2008Q2", "2009Q2"), ("2011Q2", "2013Q1"), ("2019Q3", "2020Q3")]
PERIODS = [("Pre-GFC", "2007Q2", "2008Q1"), ("GFC", "2008Q2", "2009Q2"),
           ("Pre-SDC", "2009Q3", "2011Q1"), ("SDC", "2011Q2", "2013Q1"),
           ("Pre-COVID", "2013Q2", "2019Q2"), ("COVID", "2019Q3", "2020Q3"),
           ("Post-COVID", "2020Q4", None)]
MODEL_LABEL = "IP + EUROCOIN + two PCA factors"


def read_monthly(path: Path) -> pd.DataFrame:
    frame = pd.read_csv(path, index_col="Time", parse_dates=True)
    dates = frame.index
    if (frame.empty or dates.hasnans or not dates.is_unique
            or not dates.is_monotonic_increasing or not (dates.day == 1).all()):
        raise ValueError(f"Expected distinct, ordered month-start dates: {path}")
    if not dates.equals(pd.date_range(dates[0], dates[-1], freq="MS")):
        raise ValueError(f"Missing calendar rows: {path}")
    if np.isinf(frame.to_numpy(dtype=float)).any():
        raise ValueError(f"Infinite observations: {path}")
    return frame


def target_quarters(dates: pd.DatetimeIndex) -> pd.DatetimeIndex:
    return dates.to_period("Q").asfreq("M", how="end").to_timestamp()


def error_bands(forecasts: pd.Series, growth: pd.Series, minimum: int = 40) -> pd.DataFrame:
    """Use only earlier released errors from the same month of the quarter."""
    dates = forecasts.index
    targets = target_quarters(dates)
    phase = (dates.month - 1) % 3
    errors = growth.reindex(targets).to_numpy() - forecasts.to_numpy()
    releases = targets + pd.DateOffset(months=1)
    bands = pd.DataFrame({"lower": np.nan, "upper": np.nan, "n_errors": 0}, index=dates)
    for i, date in enumerate(dates):
        past = errors[(phase == phase[i]) & (releases <= date) & np.isfinite(errors)]
        bands.loc[date, "n_errors"] = len(past)
        if len(past) >= minimum and np.isfinite(forecasts.iloc[i]):
            bands.loc[date, ["lower", "upper"]] = forecasts.iloc[i] + np.quantile(
                past, [0.05, 0.95], method="linear")
    return bands


def measures(log_growth: float, quarter: pd.Period, levels: pd.Series) -> dict:
    """Exact QoQ, YoY, level and flat-path annual carry-over conversions."""
    previous = levels.get(quarter - 1, np.nan)
    level = previous * np.exp(log_growth)
    last_year = levels.reindex(pd.period_range(f"{quarter.year-1}Q1", periods=4, freq="Q"))
    earlier = levels.reindex(pd.period_range(f"{quarter.year}Q1", periods=quarter.quarter-1, freq="Q"))
    carry = np.nan
    if last_year.notna().all() and earlier.notna().all() and np.isfinite(level):
        annual = (earlier.sum() + (5-quarter.quarter)*level) / 4
        carry = 100 * (annual / last_year.mean() - 1)
    return {"qoq_percent": 100*np.expm1(log_growth), "level_billions": level/1000,
            "yoy_percent": 100*(level/levels.get(quarter-4, np.nan)-1),
            "carry_over_percent": carry}


def score_history(history: pd.DataFrame, reference: pd.Timestamp) -> pd.DataFrame:
    """Log-growth percentage-point errors by period and month in the quarter."""
    q = history.index.to_period("Q")
    eligible = (q >= SCORE_START) & (target_quarters(history.index)+pd.DateOffset(months=1) <= reference)
    crisis = np.zeros(len(q), dtype=bool)
    for first, last in CRISES:
        crisis |= (q >= first) & (q <= last)
    masks = {"All": eligible, "Crisis": eligible & crisis, "Non-crisis": eligible & ~crisis}
    for name, first, last in PERIODS:
        masks[name] = eligible & (q >= first) & (True if last is None else q <= last)
    rows = []
    for name, mask in masks.items():
        for stage in (0, 1, 2, 3):
            use = mask & (True if stage == 0 else (history.index.month-1) % 3+1 == stage)
            values = history.loc[use].dropna(subset=["nowcast", "actual_log", "naive_log"])
            error = 100*(values.nowcast-values.actual_log)
            naive = 100*(values.naive_log-values.actual_log)
            rmse = np.sqrt(np.mean(error**2)) if len(values) else np.nan
            benchmark = np.sqrt(np.mean(naive**2)) if len(values) else np.nan
            bounded = values.dropna(subset=["lower_log", "upper_log"])
            covered = bounded.actual_log.between(bounded.lower_log, bounded.upper_log)
            rows.append(dict(period=name, stage=stage, n=len(values), rmse=rmse,
                             mae=abs(error).mean(), bias=error.mean(),
                             relative_naive_rmse=rmse/benchmark if benchmark else np.nan,
                             interval_n=len(bounded), coverage=covered.mean()))
    return pd.DataFrame(rows)


def figure(history: pd.DataFrame, actual: pd.Series, level: bool) -> go.Figure:
    """Monthly nowcasts and quarterly GDP with causal empirical error bands."""
    prefix = "level" if level else "growth"
    unit = "GDP (2020 chain-linked, billion euros)" if level else "Quarter-on-quarter GDP growth (%)"
    fig = go.Figure()
    for col, fill in [(f"{prefix}_lower", None), (f"{prefix}_upper", "tonexty")]:
        fig.add_trace(go.Scatter(x=history.index, y=history[col].tolist(), mode="lines",
                                line=dict(width=0), fill=fill, fillcolor="rgba(100,100,100,0.25)",
                                connectgaps=False, hoverinfo="skip", legendgroup="prediction_interval",
                                name="90% prediction interval", showlegend=fill is not None))
    fig.add_trace(go.Scatter(x=history.index, y=history[prefix].tolist(), name="Pseudo-real-time nowcast",
                            mode="lines", line=dict(color="#d95f0e", width=2.5),
                            connectgaps=False, hovertemplate="%{x|%b %Y}: %{y:.2f}<extra>Nowcast</extra>"))
    fig.add_trace(go.Scatter(x=actual.index, y=actual.tolist(), name="GDP realization",
                            mode="markers", marker=dict(color="#0072b2", size=6),
                            hovertemplate="%{x|%b %Y}: %{y:.2f}<extra>GDP realization</extra>"))
    fig.add_vline(x=ONLINE_START, line_dash="dash", line_color="#555555")
    fig.add_annotation(x=ONLINE_START, y=1.02, yref="paper",
                       text="First live release: Nov 2026 (reference: Oct 2026)",
                       showarrow=False, xanchor="right", font=dict(size=11, color="#555555"))
    fig.update_layout(template="plotly", autosize=True, height=510,
                      margin=dict(l=65, r=20, t=45, b=65),
                      font=dict(family="Arial, sans-serif", size=12),
                      legend=dict(orientation="h", y=-0.17, x=0.5, xanchor="center"),
                      yaxis=dict(title=unit), xaxis=dict(title="Forecast reference month",
                      range=[history.index[0], history.index[-1]+pd.DateOffset(months=2)]),
                      hovermode="x unified")
    if not level:
        fig.add_hline(y=0, line_width=1, line_dash="dash", line_color="black")
    return fig


def write_json(path: Path, values: dict) -> None:
    # Convert numpy types and missing scalars into portable JSON values.
    clean = json.loads(json.dumps(values, default=lambda x: x.item()))
    def finite(value):
        if isinstance(value, dict):
            return {key: finite(item) for key, item in value.items()}
        if isinstance(value, list):
            return [finite(item) for item in value]
        return None if isinstance(value, float) and not np.isfinite(value) else value
    path.write_text(json.dumps(finite(clean), indent=2, allow_nan=False)+"\n", encoding="utf-8")


def performance_pdf(scores: pd.DataFrame, history: pd.DataFrame, reference: pd.Timestamp, path: Path) -> None:
    """RMSE, interval coverage and COVID forecasts for the bridge equation."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

    styles = getSampleStyleSheet()
    story = [Paragraph("ITNOW: pseudo-real-time performance", styles["Title"]),
             Paragraph(MODEL_LABEL, styles["Normal"]), Spacer(1, 12)]
    table_style = TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("LINEBELOW", (0, 0), (-1, 0), .5, colors.black),
        ("LINEBELOW", (0, -1), (-1, -1), .5, colors.black), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)])
    captions = [
        ("Table 1. RMSE by forecast month", ["All", "Crisis", "Non-crisis", "Post-COVID"], "rmse"),
        ("Table 2. RMSE relative to the last-quarter GDP benchmark", ["All", "Crisis", "Non-crisis", "Post-COVID"], "relative_naive_rmse"),
        ("Table 3. RMSE by subperiod", [p[0] for p in PERIODS], "rmse"),
    ]
    for title, periods, metric in captions:
        story += [Paragraph(title, styles["Heading3"])]
        rows = [["Period", "Month 1", "Month 2", "Month 3", "Pooled"]]
        for period in periods:
            part = scores[scores.period == period].set_index("stage")
            rows.append([period]+[f"{part.loc[s,metric]:.3f}" for s in (1, 2, 3, 0)])
        table = Table(rows, colWidths=[135, 78, 78, 78, 78])
        table.setStyle(table_style)
        story += [table, Spacer(1, 8)]
    end = reference.to_period("Q")-1
    note = (f"Sample: Q2 2007 through {end}. RMSE is measured in log-growth percentage points. "
            "Crisis periods: Q2 2008-Q2 2009, Q2 2011-Q1 2013, Q3 2019-Q3 2020. "
            "Pooled scores contain all three monthly forecasts per quarter. "
            "We use revised data and fixed publication delays.")
    story += [Paragraph(note, styles["Normal"])]
    story += [PageBreak(), Paragraph("ITNOW: intervals and COVID forecasts", styles["Title"]),
              Paragraph("Table 4. Empirical 90% interval coverage", styles["Heading3"])]
    rows = [["Month within quarter", "Scored intervals", "Coverage (%)"]]
    for stage, label in [(1, "Month 1"), (2, "Month 2"), (3, "Month 3"), (0, "Pooled")]:
        row = scores[(scores.period == "All") & (scores.stage == stage)].iloc[0]
        coverage = f"{100*row.coverage:.1f}" if row.interval_n else "n.a."
        rows.append([label, str(int(row.interval_n)), coverage])
    table = Table(rows, colWidths=[200, 120, 120]); table.setStyle(table_style)
    story += [table, Spacer(1, 8), Paragraph(
        "Intervals use only earlier released errors from the same month within each quarter, "
        "with at least 40 errors.", styles["Normal"]),
        Spacer(1, 12), Paragraph("Table 5. COVID nowcasts and realizations", styles["Heading3"])]
    rows = [["Reference month", "GDP quarter", "Nowcast (%)", "GDP (%)"]]
    for date, row in history.loc["2020-04-01":"2020-09-01"].iterrows():
        rows.append([date.strftime("%b %Y"), str(date.to_period("Q")),
                     f"{100*np.expm1(row.nowcast):.2f}", f"{100*np.expm1(row.actual_log):.2f}"])
    table = Table(rows, colWidths=[130, 105, 105, 105]); table.setStyle(table_style)
    story += [table, Spacer(1, 8), Paragraph(
        "Quarter-on-quarter percentage changes.", styles["Normal"])]
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=45, rightMargin=45,
                      topMargin=35, bottomMargin=35).build(story)


def live_history(base: Path, website: Path, latest: dict,
                 release_date: pd.Timestamp) -> tuple[Path, bytes] | None:
    """Prepare one current monthly release without changing earlier records."""
    released = pd.Timestamp(release_date)
    if released.tzinfo is not None:
        released = released.tz_convert("Europe/Rome").tz_localize(None)
    released = released.normalize()
    reference = pd.Timestamp(latest["reference_date"])
    produced = pd.Timestamp(latest["produced_at"]).tz_convert("Europe/Rome").tz_localize(None)
    vintage = pd.Timestamp(latest["data_vintage"])
    month = released.to_period("M")
    if (released < FIRST_LIVE_RELEASE or reference != (month-1).to_timestamp()
            or produced.to_period("M") != month
            or produced.normalize() > released or vintage > produced.normalize()):
        return None

    from openpyxl import load_workbook

    local = base/"HISTORY/RELEASES/ITNOW_History.xlsx"
    public = website/"static/ITNOW/ITNOW_History.xlsx"
    template = local if local.is_file() else public
    if not template.is_file():
        raise FileNotFoundError(f"Supply the release workbook at {local} or {public}.")
    original = template.read_bytes()
    workbook = load_workbook(BytesIO(original))
    headers = ["Release date", "Reference month", "Target quarter", "QoQ", "Lower 5%",
               "Upper 95%", "YoY", "Carry-over", "GDP level (EUR bn)", "Data vintage",
               "Produced at", "Code version"]

    def records(book):
        if not {"INFO", "DATA", "NOWCAST"}.issubset(book.sheetnames):
            raise ValueError("The release workbook needs INFO, DATA and NOWCAST sheets.")
        sheet = book["NOWCAST"]
        if [cell.value for cell in sheet[1]] != headers:
            raise ValueError("Unexpected NOWCAST columns in the release workbook.")
        return [row for row in sheet.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in row)]

    existing = records(workbook)
    if public.is_file():
        published = records(load_workbook(BytesIO(public.read_bytes())))
        if any(row not in existing for row in published):
            raise ValueError("Local and website release histories differ; synchronize them before adding a release.")
    current = latest["current"]
    rate_keys = ["qoq_percent", "lower_percent", "upper_percent", "yoy_percent", "carry_over_percent"]
    rates = [float(current[key])/100 if current[key] is not None and np.isfinite(current[key])
             else None for key in rate_keys]
    level = current["level_billions"]
    if level is not None and not np.isfinite(level):
        level = None
    values = [released.to_pydatetime(), reference.to_pydatetime(), latest["quarter"], *rates,
              level, vintage.to_pydatetime(), latest["produced_at"],
              latest.get("code_version")]
    matches = [row for row in existing if pd.Timestamp(row[1]) == reference
               and pd.Timestamp(row[9]) == vintage
               and row[10] == latest["produced_at"] and row[11] == latest.get("code_version")]
    if matches:
        same = len(matches) == 1
        for before, after in zip(matches[0][1:], values[1:]):
            if isinstance(before, (int, float)) and isinstance(after, (int, float)):
                same &= bool(np.isclose(before, after, rtol=1e-12, atol=1e-14))
            else:
                same &= before == after
        if not same:
            raise ValueError(f"A different release for {reference:%B %Y} is already recorded; it will not be overwritten.")
        return local, original

    sheet = workbook["NOWCAST"]
    target = max((row[0].row for row in sheet.iter_rows(min_row=2)
                  if any(cell.value is not None for cell in row)), default=1)+1
    for column, value in enumerate(values, 1):
        cell = sheet.cell(target, column)
        cell._style = copy(sheet.cell(max(2, target-1), column)._style)
        cell.value = value
    buffer = BytesIO()
    workbook.save(buffer)
    return local, buffer.getvalue()


def stage_website(root: Path, assets: Path, project: Path,
                  release: tuple[Path, bytes] | None = None) -> None:
    """Copy only ITNOW sources and assets; keep other website work untouched."""
    if not (root/"content").is_dir() or not (root/"static").is_dir():
        raise ValueError("website-root must be the existing local website checkout.")
    for name in ("MODEL.pdf", "GDP_MEASURES.pdf"):
        if not (assets/name).is_file() and not (root/"static/ITNOW"/name).is_file():
            raise FileNotFoundError(f"Supply {name} in the report or website assets before staging.")
    backup = assets/"website_backup"
    for source in sorted((project/"website").rglob("*")):
        if not source.is_file():
            continue
        relative = source.relative_to(project/"website")
        target = root/relative
        if target.exists():
            saved = backup/relative
            if not saved.exists():
                saved.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(target, saved)
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
    destination = root/"static/ITNOW"
    destination.mkdir(parents=True, exist_ok=True)
    for source in assets.iterdir():
        if source.is_file() and source.name not in {
                "release_record.csv", "ITNOW_History.xlsx", "ITNOW_History_pre_bridge.xlsx"}:
            saved = backup/"static/ITNOW"/source.name
            target = destination/source.name
            if target.exists() and not saved.exists():
                saved.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(target, saved)
            shutil.copy2(source, target)
    if release is not None:
        local, content = release
        targets = [local, destination/"ITNOW_History.xlsx"]
        previous = {path: path.read_bytes() if path.is_file() else None for path in targets}
        try:
            for path in targets:
                path.parent.mkdir(parents=True, exist_ok=True)
                path.write_bytes(content)
        except Exception:
            for path, old in previous.items():
                if old is None:
                    path.unlink(missing_ok=True)
                else:
                    path.write_bytes(old)
            raise


def build_report(base: Path, vintage: str, website: Path | None = None) -> Path:
    project = Path(__file__).resolve().parents[1]
    prediction = base/"HISTORY/PREDICTION"/f"{vintage}_bridge.csv"
    predicted = read_monthly(prediction)
    gdp_path = base/"DATA/GDP_History"/f"GDP{vintage}.csv"
    gdp = read_monthly(gdp_path)
    metadata = json.loads(prediction.with_suffix(".json").read_text())
    source = json.loads((base/"DATA/VINTAGE_METADATA"/f"vintage{vintage}.json").read_text())
    reference = predicted.index[-1]
    produced = pd.Timestamp(metadata["produced_at"])
    if (metadata["vintage"] != vintage or pd.Timestamp(metadata["eval_end"]) != reference
            or metadata["model"] != "ip_eurocoin_pca_2" or produced.tzinfo is None):
        raise ValueError("Forecast metadata disagrees with the selected file.")
    if set(predicted) != {"nowcast"} or not np.isfinite(predicted.nowcast.iloc[-1]):
        raise ValueError("The selected model must have a finite current nowcast.")
    if (gdp.Growth.notna() & (gdp.index.month % 3 != 0)).any():
        raise ValueError("GDP must be dated to quarter-end months.")
    # Prevent unreleased outcomes in an intentionally older report.
    gdp.loc[gdp.index+pd.DateOffset(months=1) > reference, :] = np.nan
    levels = gdp.Level.dropna()
    levels.index = levels.index.to_period("Q")
    dates = predicted.index
    targets = target_quarters(dates)
    quarters = dates.to_period("Q")
    previous = levels.reindex(quarters-1).to_numpy()
    bands = error_bands(predicted.nowcast, gdp.Growth)
    history = predicted.copy()
    history["actual_log"] = gdp.Growth.reindex(targets).to_numpy()
    history["naive_log"] = gdp.Growth.reindex(targets-pd.DateOffset(months=3)).to_numpy()
    history["lower_log"], history["upper_log"] = bands.lower, bands.upper
    history["n_errors"] = bands.n_errors
    for suffix, logs in [("", history.nowcast), ("_lower", bands.lower), ("_upper", bands.upper)]:
        history["growth"+suffix] = 100*np.expm1(logs)
        history["level"+suffix] = previous*np.exp(logs)/1000
    stats = measures(history.nowcast.iloc[-1], quarters[-1], levels)
    stats.update(lower_percent=history.growth_lower.iloc[-1], upper_percent=history.growth_upper.iloc[-1],
                 previous_qoq_percent=100*np.expm1(history.naive_log.iloc[-1]))
    scores = score_history(history, reference)
    output = base/"output"/vintage
    assets = output/"website/ITNOW"
    assets.mkdir(parents=True, exist_ok=True)
    history.to_csv(assets/"Retrospective_Nowcasts.csv", index_label="Time", date_format="%Y-%m-%d")
    scores.to_csv(output/"performance.csv", index=False)
    for level, name in [(False, "GDP_Growth_Ncast.json"), (True, "GDP_Level_Ncast.json")]:
        actual = gdp.Level/1000 if level else 100*np.expm1(gdp.Growth)
        actual = actual.loc[dates[0]:reference].dropna()
        figure(history, actual, level).write_json(assets/name)
    (assets/"plotly.min.js").write_text(get_plotlyjs(), encoding="utf-8")
    last_actual = str(levels.index[-1])
    latest = dict(model=MODEL_LABEL, model_id=metadata["model"], model_version=metadata["model_version"],
                  code_version=metadata.get("code_version"),
                  reference_month=reference.strftime("%B %Y"), reference_date=reference.date().isoformat(),
                  quarter=f"Q{quarters[-1].quarter} {quarters[-1].year}", produced_at=metadata["produced_at"],
                  production_date=produced.tz_convert("Europe/Rome").date().isoformat(),
                  data_vintage=vintage, data_date=pd.Timestamp(vintage).strftime("%d %B %Y"),
                  current=stats, last_actual_quarter=last_actual, gdp_source=source["gdp"],
                  interval_n=int(bands.n_errors.iloc[-1]),
                  band_coverage=scores[(scores.period == "All") & (scores.stage == 0)].iloc[0].coverage,
                  input_hashes={str(path.relative_to(base)): hashlib.sha256(path.read_bytes()).hexdigest()
                                for path in [prediction, gdp_path]},
                  series_kind="retrospective calculation under the current model",
                  first_live_release_month=FIRST_LIVE_RELEASE.date().isoformat(),
                  first_live_reference_month=ONLINE_START.date().isoformat())
    write_json(assets/"latest_itnow.json", latest)
    # Save the current record for publication review.
    pd.DataFrame([{**metadata, **stats, "status": "unpublished"}]).to_csv(assets/"release_record.csv", index=False)
    performance_pdf(scores, history, reference, assets/"ITNOW_Performance.pdf")
    (base/"output/pdf").mkdir(parents=True, exist_ok=True)
    shutil.copy2(assets/"ITNOW_Performance.pdf", base/"output/pdf/ITNOW_Performance.pdf")
    for name in ("MODEL", "GDP_MEASURES"):
        pdf = base/"output/pdf"/f"{name}.pdf"
        if pdf.is_file():
            shutil.copy2(pdf, assets/pdf.name)
    if website:
        release = live_history(base, website, latest, pd.Timestamp.now(tz="Europe/Rome"))
        stage_website(website, assets, project, release)
    print(f"Report: {assets}\n{latest['quarter']}: {stats['qoq_percent']:.3f}% QoQ; reference {latest['reference_month']}.")
    return assets


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-path", type=Path, required=True)
    parser.add_argument("--vintage", required=True)
    parser.add_argument("--website-root", type=Path)
    args = parser.parse_args()
    build_report(args.base_path.resolve(), args.vintage, args.website_root)
